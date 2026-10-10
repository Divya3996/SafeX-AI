#!/usr/bin/env python3
"""Fill the supplied Code Carnival template without changing event branding.

Dependencies: python-pptx, Pillow, CairoSVG, qrcode. PDF export uses LibreOffice.
Run from the project root. Sources are local release records, not inferred claims.
"""
from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
import json
import re

from PIL import Image
import cairosvg
import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentations"
ASSETS = OUT / "assets"
TEMPLATE = Path("/home/viyat/Downloads/Code Carnival 3.0 Idea Submission Template.pptx")
TEAM, TEAM_ID = "Doclock", "M87V"
REPO = "https://github.com/Divya3996/SafeX-AI"
DEMO = "https://safex-ai-jury-demo.cuddlyrobin1.chatgpt.site"
STEM = f"{TEAM_ID}_{TEAM}_CC3"
RED, INK, CREAM, MUTED, WHITE = "C8102E", "1E1917", "F6F0E6", "5A504A", "FFFFFF"
PUBLIC, OSWALD = "Public Sans", "Oswald"

TIMES = [35, 55, 55, 90, 100, 60, 75, 65, 65]
SPEAKERS = ["Divya Gohil"] * 5 + ["Krupa"] * 4
SCRIPT = [
    "Hello, we are team Doclock, M87V. I am Divya Gohil, the team leader, and "
    "Krupa is my teammate. Our project is SafeX AI: an on-device security assistant "
    "for phishing links, scam messages and suspicious communications. We have a "
    "working Android app and Chrome extension. Our goal is to help people pause, "
    "understand the evidence and choose a safer next action, without uploading "
    "their private messages or screenshots for AI inference.",

    "The problem happens at the moment someone is asked to trust a message. A "
    "message may imitate a bank, promise a reward, demand an urgent payment or ask "
    "for an OTP. Those requests arrive through different apps, links and QR codes. "
    "A user needs understandable help before acting. We also considered local "
    "languages, comfortable reading sizes and poor connectivity. Some scams spread "
    "across several messages: an ordinary-looking reward promise may become "
    "suspicious only when a later message demands money to release it. Our design "
    "therefore combines individual checks with explicit, user-reviewed context.",

    "SafeX AI combines compact trained classifiers with security rules. The "
    "Android app accepts deliberate inputs and optional readable notifications; "
    "the Chrome extension adds authorized page and link context. The output "
    "includes the warning, reasons, coverage limits and next steps. We do not "
    "call a cloud chatbot. For screenshots and QR codes, local recognition "
    "prepares content the user can review. Scam Story Mode adds an evidence "
    "timeline with exact supporting passages. The interface and guidance are "
    "available in English, Hindi and Gujarati, with 85 to 150 percent text scaling.",

    "There are four feature groups. First, Android can check text, links, "
    "screenshots, QR images, live camera QR and supported file structures. "
    "Notification analysis and warning sound are optional. Second, the floating "
    "shield lets the user give screen-capture consent, crop an area, correct "
    "recognized text and analyze privately. It does not silently capture other "
    "apps. Third, Chrome checks selected text, real link destinations and "
    "authorized pages. Optional site protection adds bounded rescans and "
    "warnings on supported risky clicks and form submissions. Fourth, Scam "
    "Story Mode connects reviewed evidence, while tours and incident-help "
    "guides make the product easier to use. Android has a 19-step highlighted "
    "tour and Chrome has 12 steps. Show one selected-message scan and its reasons "
    "here; use the longer floating or browser demonstration only if time permits.",

    "Every input follows the same idea: choose content, extract and review it "
    "when necessary, analyze locally, then explain the result. Text uses a "
    "logistic classifier and request-context rules; URLs use a structural neural "
    "network and URL rules. Chrome adds visible link and form metadata while "
    "excluding field values. In Android Scam Story Mode, add the two authored "
    "messages 'Your commission balance is 5000 rupees' and 'Transfer 1000 rupees "
    "to release the balance'. The installed fixture records no individual "
    "warning for either message, but the supported sequence connects the "
    "balance promise to the payment demand. Open the finding to show its exact "
    "supporting passages. Remove the payment item and the connection disappears. "
    "That is one controlled development example, not a general accuracy claim. "
    "Keep airplane mode enabled and do not open a fraud link or make a payment.",

    "The app uses Kotlin, Compose, Hilt and Room; Chrome uses TypeScript and "
    "Manifest V3. The message model is logistic regression with word, bigram "
    "and character-trigram features. Its 5606 prepared messages are split into "
    "3910 training, 810 validation and 886 evaluation examples. Observed "
    "English data is supplemented with authored Hindi and Gujarati examples. "
    "The URL neural network has 15 inputs and layers of 24, 12 and one output, "
    "exported to TensorFlow Lite/LiteRT. Its sources include PhiUSIIL, authored "
    "variants and a bounded reported-link dataset. ML Kit and Tesseract provide "
    "bundled OCR/QR recognition. Models are trained on the development computer "
    "and packaged for local inference; the phone does not automatically retrain "
    "on users' messages. Android requests no Internet permission.",

    "We have installable Android 1.9.0 and an unpacked Chrome 1.1.0 release. "
    "The final Android build passed 593 active unit tests and 73 active offline "
    "Android checks, with zero lint errors and 108 warnings. Its tests used "
    "Android 14 on an emulator, not our Samsung phones. Chrome's recorded "
    "validation passed 431 unit checks and 43 browser scenarios. These counts "
    "establish functional behavior. Separately, the Android 1.7 historical "
    "English development comparison detected 71 of 78 scams, with one warning "
    "among 769 legitimate messages; it still missed seven scams. The models "
    "are unchanged in 1.9. We do not present that result as universal accuracy "
    "or as a browser benchmark. Physical S24/A36 testing, fresh independent "
    "multilingual evaluation and publisher release work remain open.",

    "Our immediate goal is a reliable end-to-end rehearsal on both Samsung "
    "phones and desktop Chrome, including capture consent, OCR review, sound "
    "and permission recovery. After the hackathon we plan a consented pilot "
    "and fresh independently labeled scam and legitimate examples. We will "
    "measure recall and false warnings by language and scam category before "
    "adding broader sequence coverage. Longer-term work includes signed "
    "publisher-controlled model updates, wider compatibility, store releases "
    "and privacy-preserving education or institutional pilots. These are "
    "planned steps, not shipped capabilities or existing partnerships.",

    "Our public repository is github.com/Divya3996/SafeX-AI. The source "
    "documentation covers setup, features, privacy and validation. The QR code "
    "opens our public jury demonstration library with labeled messages, "
    "URLs, QR codes and screenshots. The public website has older versioned "
    "downloads, so use the separately supplied Android 1.9.0 APK and Chrome "
    "1.1.0 ZIP for this presentation. Our goal is security guidance that people "
    "can understand, use offline and read in their own language. The floating "
    "assistant supports phone workflows, Chrome adds browser context, and "
    "Scam Story Mode connects the clues users deliberately review. The next "
    "step is independent multilingual evaluation and physical-device rehearsal. "
    "Thank you; we welcome your questions.",
]

DETAILS = [
    "Team name and ID are user supplied. Names supplied: Divya Gohil (leader), "
    "Krupa (member). 'On-device AI Security' is a descriptive theme based on the "
    "selected statement, not a verified organizer track code. No invented surname, "
    "contact information, registration value or submission deadline is used.",
    "No prevalence, financial-loss or adoption statistic is invented. Intended "
    "users and usability benefits are design goals, not a completed field study.",
    "Android ALLOW/WARN/BLOCK describes the local decision/opening policy. "
    "A model alone can add a warning, not independently block. A risk index is "
    "not a calibrated probability. A missing pattern/list entry does not certify safety.",
    "Android also has explicit share/selected-text input, UPI address/amount "
    "comparison, disclosed context questions, local searchable history with "
    "7/30/90-day retention, warning-sound test, and country-aware incident "
    "guidance. File inspection is bounded structure checking, not antivirus. "
    "Chrome also has pointer crop, local OCR/QR, redacted reports/export, "
    "incorrect-result feedback, themes, keyboard support, warning sound and "
    "opt-in exact URL blocking from explicitly labeled local list entries. "
    "Scam Story Mode is currently Android-only. Neither product continuously "
    "reads the clipboard, authenticates callers or intercepts all communications.",
    "The Android sequence engine has three bounded playbooks: task/release "
    "fee, claimed authority plus pressure/demand, and support plus remote "
    "access and sensitive requests. Cases accept up to eight reviewed items. "
    "Individual warnings are retained. Cases without a sequence match say "
    "more context is needed. Drafts live in app-owned RAM, clear on screen "
    "lock and expire after two minutes away. Explicit case snapshots use "
    "AES-256-GCM with an Android Keystore key; exports are previewed and "
    "exclude raw contents, URLs and payment details. Ordinary Android "
    "history and Chrome report storage do not have that case-vault encryption guarantee.",
    "Text feature vocabulary: 13605 terms; shipped JSON about 462 KB. URL "
    "model: 15 -> 24 -> 12 -> 1, 4016-byte TFLite asset. The research text "
    "threshold is 0.545 with action context; a guarded familiar-Latin path "
    "uses 0.915. URL warning threshold is about 0.9001. These are internal "
    "uncalibrated scores, not fraud probabilities. Campaign grouping and "
    "domain separation reduce split overlap. Model loading checks hashes, "
    "feature order and finite values; failure retains local rules. Training "
    "and model artifacts were promoted in 1.7 and not retrained for Story Mode.",
    "Four optional unit research-export checks and one optional native "
    "research-input check were skipped; the 593/73 numbers count passed "
    "active checks. Twenty authored Story contracts matched, including ten "
    "suspicious and ten legitimate comparisons. One paid-support control "
    "intentionally gets a review connection, while existing per-item "
    "warnings remain. This is not 100 percent detection accuracy. Text "
    "model-only historical English recall was 76/78 (97.4 percent); full "
    "app recall was 71/78 (91.0 percent). URL model-only historical recall "
    "was 93.31 percent with 0.40 percent false warnings; new-domain "
    "reported-feed recall was 68.02 percent, with no matched current benign "
    "set. Hindi/Gujarati independent real-world accuracy is not established. "
    "Shown warning screenshots are existing prior-release captures, cropped "
    "to the relevant region, not a claim of new 1.9 screenshots.",
    "Physical S24/A36, Android 16/16 KB runtime, TalkBack/manual accessibility "
    "and physical speaker checks remain. Publisher signing and store "
    "reviews remain. No dates, pilot participants, clinical/financial "
    "outcomes, active partnerships or automatic retraining are promised. "
    "Fresh evaluation must be kept separate from implementation and tuning.",
    "Public GitHub API verified private=false and an accessible README on "
    "10 October 2026. Contributor listing contained Divya3996 only. Current "
    "1.9 changes exist locally and are not pushed by this presentation "
    "task. The public jury website returned HTTP 200 and SafeX-branded "
    "HTML on 10 October. Its library has 271 development fixtures; do not "
    "claim these are independent tests. Do not use its older app/extension "
    "version references as current release references.",
]

SOURCES = [
    ["User-provided team details", "Supplied Code Carnival 3.0 template"],
    ["docs/features.md", "docs/mentor-meeting-plan.md (historical design context)"],
    ["README.md", "docs/architecture.md", "docs/scam-story-mode.md"],
    ["docs/features.md", "docs/practical-features.md", "docs/onboarding-and-release.md",
     "chrome-extension/README.md", "docs/chrome-extension-release.md"],
    ["docs/architecture.md", "docs/scam-story-mode.md",
     "docs/test-results/story-evaluation.json", "docs/privacy.md"],
    ["docs/ml-model.md", "models/fraud-text-evaluation.json",
     "models/fraud-url-evaluation.json", "models/FRAUD-DATA-ATTRIBUTION.md",
     "android-app/app/src/main/assets/text-model-card.json",
     "android-app/app/src/main/assets/url-model-card.json"],
    ["docs/test-results/story-release-summary.json",
     "docs/test-results/chrome-extension-summary.json",
     "docs/test-results/fraud-native-evaluation.json", "docs/validation.md",
     "docs/screenshots/scam-result.png", "docs/screenshots/chrome-extension-scan-warning.png"],
    ["docs/validation.md", "docs/chrome-extension-release.md", "docs/ml-model.md"],
    [REPO, DEMO, "docs/jury-demo.md", "README.md"],
]


def remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def write(shape, value, *, size=None, height=None, width=None, x=None, y=None):
    """Preserve template run/paragraph styles while replacing content."""
    tf = shape.text_frame
    old = tf.paragraphs[0]
    rpr = deepcopy(old.runs[0]._r.rPr) if old.runs and old.runs[0]._r.rPr is not None else None
    ppr = deepcopy(old._p.pPr) if old._p.pPr is not None else None
    tf.clear()
    tf.word_wrap = True
    for index, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        if ppr is not None:
            p._p.insert(0, deepcopy(ppr))
        run = p.add_run()
        run.text = line
        if rpr is not None:
            run._r.insert(0, deepcopy(rpr))
        if size is not None:
            run.font.size = Pt(size)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.1
    for attr, val in [("height", height), ("width", width), ("left", x), ("top", y)]:
        if val is not None:
            setattr(shape, attr, Inches(val))
    return shape


def bullets(shape, lines, *, y=None, height=2.95, size=18, gap=12):
    write(shape, "\n".join(lines), size=size, height=height, y=y)
    for p in shape.text_frame.paragraphs:
        pp = p._p.get_or_add_pPr()
        for child in list(pp):
            if child.tag.endswith(("buNone", "buChar", "buAutoNum")):
                pp.remove(child)
        pp.set("marL", str(Inches(.16)))
        pp.set("indent", str(-Inches(.16)))
        dot = OxmlElement("a:buChar")
        dot.set("char", "•")
        pp.append(dot)
        p.line_spacing = 1.15
        p.space_after = Pt(gap)


def text(slide, value, x, y, w, h, *, size=18, bold=False, color=INK,
         font=PUBLIC, align=PP_ALIGN.LEFT, name=None, link=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.08
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        run.font.name, run.font.size = font, Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        if link:
            run.hyperlink.address = link
    return shape


def box(slide, x, y, w, h, *, fill=WHITE, stroke=RED, name=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = .05
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.color.rgb = RGBColor.from_string(stroke)
    shape.line.width = Pt(1)
    if name:
        shape.name = name
    return shape


def arrow(slide, x, y, w=.3, h=.16):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(RED)
    shape.line.fill.background()
    return shape


def picture(slide, path, x, y, w, h):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def main():
    OUT.mkdir(exist_ok=True)
    ASSETS.mkdir(exist_ok=True)
    r = Presentation(TEMPLATE)
    original_titles = [r.slides[i].shapes[5 if i != 0 and i != 9 else 6].text
                       for i in range(10) if i != 1]
    # Snapshot all original organizer pictures; neither pixels nor geometry may change.
    images = []
    for i, s in enumerate(r.slides):
        if i == 1:
            continue
        for sh in s.shapes:
            if sh.shape_type == 13:
                images.append({"slide": i, "name": sh.name,
                               "sha256": sha256(sh.image.blob).hexdigest(),
                               "geometry": [sh.left, sh.top, sh.width, sh.height]})
    # Resolve originals before any removal changes shape indices.
    shapes = [list(s.shapes) for s in r.slides]
    slides = list(r.slides)

    # 1. Event title and organizer branding remain intact.
    s, sh = slides[0], shapes[0]
    write(sh[7], "SafeX AI · Android + Chrome · Private, on-device protection", size=18, height=.36)
    for idx, value in [(10, TEAM), (13, TEAM_ID), (16, "Divya Gohil"),
                       (19, "On-device threat, phishing and scam detection"),
                       (22, "On-device AI Security")]:
        write(sh[idx], value, size=18, height=.37)
    text(s, "Team member: Krupa", .89, 1.82, 6.3, .36, color=CREAM, name="Team member")
    cairosvg.svg2png(url=str(ROOT / "branding/safex-logo.svg"),
                    write_to=str(ASSETS / "safex-logo.png"), output_width=512, output_height=512)
    picture(s, ASSETS / "safex-logo.png", 10.57, 1.7, .53, .53)
    text(s, "SafeX AI", 11.23, 1.8, 1.2, .35, color=CREAM, bold=True, name="Product wordmark")

    # 2. Problem: exactly six concise bullets across the supplied two panels.
    sh = shapes[2]
    bullets(sh[9], [
        "Phishing links, payment traps and impersonation arrive across apps.",
        "People must decide before sharing secrets, installing apps or paying.",
        "Local-language readers need clear, accessible guidance.",
    ])
    bullets(sh[12], [
        "Urgency and trusted branding can hide a dangerous request.",
        "Single-message checks can miss a scam spread across messages.",
        "Privacy, offline use and understandable evidence must work together.",
    ])

    # 3. Solution: four bullets and an editable concept diagram.
    s, sh = slides[3], shapes[3]
    bullets(sh[9], [
        "Local AI + rules explain suspicious requests.",
        "Android app and Chrome extension fit everyday workflows.",
        "Scam Story Mode connects reviewed evidence.",
        "English / Hindi / Gujarati; larger text and guided setup.",
    ], gap=9)
    for shape in (sh[11], sh[12]):
        remove_shape(shape)
    text(s, "FROM CONTENT TO A CLEAR NEXT STEP", 7.06, 2.77, 5.03, .4,
         size=18, bold=True, color=RED, font=OSWALD)
    for y, value in [(3.29, "Message / link / screenshot / QR"),
                     (4.09, "On-device models + security rules"),
                     (4.89, "Warning + evidence + next step")]:
        box(s, 7.15, y, 4.95, .58)
        text(s, value, 7.3, y+.12, 4.64, .34, align=PP_ALIGN.CENTER)
    for y in (3.91, 4.71):
        a = arrow(s, 9.52, y, .21, .14)
        a.rotation = 90
    text(s, "User-controlled. No cloud inference.", 7.08, 5.85, 5.02, .4,
         size=18, color=MUTED, align=PP_ALIGN.CENTER)

    # 4. Four core groups, following the supplied feature grid.
    sh = shapes[4]
    features = [
        (9, 10, "Android scanning", "Scan messages, links, screenshots, QR and files. Optional notification warnings and sound."),
        (13, 14, "Floating assistant", "Shield → consent → crop → edit → scan. Paste or import an image for a private result."),
        (17, 18, "Chrome protection", "Review text, links and pages. Optional site monitoring and risky click/form warnings."),
        (21, 22, "Scam Story + accessible UX", "Connect messages; save encrypted cases. Three languages, larger text, tours and scam help."),
    ]
    for title_idx, desc_idx, title, desc in features:
        write(sh[title_idx], title, size=18, height=.35)
        write(sh[desc_idx], desc, size=18, height=1.12)

    # 5. Existing four-stage workflow plus editable local architecture diagram.
    s, sh = slides[5], shapes[5]
    for idx, label in [(9, "Choose content"), (13, "Extract & review"),
                       (17, "Analyze locally"), (21, "Explain & act")]:
        write(sh[idx], label, size=18, height=.34)
    for shape in (sh[23], sh[24]):
        remove_shape(shape)
    nodes = [(1.12, 2.54, "REVIEWED INPUT", "Text / links\nLocal OCR + QR"),
             (4.09, 4.65, "LOCAL SECURITY ENGINE", "Text + URL models; rules\nBrowser context; story clues"),
             (9.17, 2.99, "USER OUTPUT", "Warning + reasons\nNext-step guidance")]
    for x, w, title, body in nodes:
        box(s, x, 3.83, w, 1.53)
        text(s, title, x+.15, 4.00, w-.3, .35, bold=True, color=RED, font=OSWALD)
        text(s, body, x+.15, 4.49, w-.3, .7)
    arrow(s, 3.72, 4.53, .25, .16)
    arrow(s, 8.85, 4.53, .25, .16)
    text(s, "Authored story: balance promise → release payment → connected warning",
         1.13, 5.47, 10.96, .4, size=18, bold=True, color=RED)
    text(s, "Original passages support the finding; the user reviews every case item.",
         1.13, 5.95, 10.96, .35, size=18, color=MUTED)

    # 6. Keep supplied table and its colors; actual technology only.
    sh = shapes[6]
    table = sh[10].table
    rows = [
        ["Layer", "Technology", "Why we chose it"],
        ["Frontend", "Kotlin / Compose; TypeScript", "Native app + browser interface"],
        ["Runtime", "Android + Chrome MV3", "On-device inference; no cloud backend"],
        ["Storage", "Room; AES-GCM; Chrome local", "History, encrypted cases, reports"],
        ["AI / recognition", "Logistic model; LiteRT; OCR", "Offline text, URL, image checks"],
        ["Hosting", "Public static demo website", "Shareable jury fixtures and guide"],
        ["Build / testing", "Gradle; Python; Playwright", "Reproducible builds and validation"],
    ]
    for row_idx, values in enumerate(rows):
        table.rows[row_idx].height = Inches(.48)
        for col_idx, val in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            old = cell.text_frame.paragraphs[0]
            rpr = deepcopy(old.runs[0]._r.rPr) if old.runs else None
            cell.text_frame.clear()
            cell.text_frame.word_wrap = True
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = val
            if rpr is not None:
                run._r.insert(0, rpr)
            run.font.size = Pt(18)
            p.line_spacing = 1.05
            p.space_before = p.space_after = Pt(0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    sh[10].height = Inches(.48*7)
    sh[7].height = Inches(.48*7)
    sh[8].height = Inches(.48*7)
    sh[9].top = Inches(2.53+.48*7)
    text(slides[6], "Bundled models and OCR/QR assets; no cloud inference API.",
         1.02, 6.10, 11.18, .4, size=18, color=MUTED)

    # 7. Working progress versus remaining gates, six bullets total.
    s, sh = slides[7], shapes[7]
    bullets(sh[9], ["Android 1.9.0 + Chrome 1.1.0.",
                    "593 Android unit + 73 offline checks passed.",
                    "431 Chrome unit + 43 browser checks passed."], height=2.85, gap=13)
    bullets(sh[12], ["S24/A36 rehearsal and accessibility review.",
                     "Independent multilingual fraud evaluation.",
                     "Publisher signing and store submission."], height=2.85, gap=13)
    for shape in (sh[14], sh[15]):
        remove_shape(shape)
    text(s, "ACTUAL PRODUCT CAPTURES", 8.98, 2.77, 3.22, .35,
         size=18, bold=True, color=RED, font=OSWALD)
    im = Image.open(ROOT / "docs/screenshots/scam-result.png")
    im.crop((35, 600, 685, 810)).save(ASSETS / "android-warning-detail.png")
    im = Image.open(ROOT / "docs/screenshots/chrome-extension-scan-warning.png")
    im.crop((25, 1025, 297, 1114)).save(ASSETS / "chrome-warning-detail.png")
    text(s, "Android warning", 9.00, 3.25, 3.2, .34, size=18)
    picture(s, ASSETS / "android-warning-detail.png", 9.00, 3.65, 3.15, 1.018)
    text(s, "Chrome warning", 9.00, 4.86, 3.2, .34, size=18)
    picture(s, ASSETS / "chrome-warning-detail.png", 9.00, 5.24, 3.15, 1.031)
    # Screenshot provenance is fully documented in notes and the companion guide.

    # 8. Six concrete planned actions across the provided three phases.
    sh = shapes[8]
    bullets(sh[10], [
        "Rehearse app + extension on S24/A36; verify sound and consent.",
        "Show scam/control pairs, connected evidence and offline scans.",
    ], height=2.7, gap=15)
    bullets(sh[14], [
        "Collect fresh labeled EN/HI/GU examples; evaluate by scam category.",
        "Run a consented pilot; refine warnings and add story playbooks.",
    ], height=2.7, gap=15)
    bullets(sh[18], [
        "Signed model updates, wider device support and store releases.",
        "Explore bank/education partnerships through private pilots.",
    ], height=2.7, gap=15)

    # 9. Real verified repository and demo; no invented contribution claims.
    s, sh = slides[9], shapes[9]
    write(sh[10], REPO, size=24, height=.5)
    sh[10].click_action.hyperlink.address = REPO
    for idx, value in [(13, "Public access verified.\nNo sign-in required."),
                       (16, "Setup, features, privacy and validation documented."),
                       (19, "Development history available. Final source sync pending."),
                       (22, "Public jury demo")]:
        write(sh[idx], value, size=18, height=1.85 if idx != 22 else .4)
    sh[22].click_action.hyperlink.address = DEMO
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M,
                       box_size=10, border=4)
    qr.add_data(DEMO)
    qr.make(fit=True)
    qr.make_image(fill_color="black", back_color="white").save(ASSETS / "public-demo-qr.png")
    pic = picture(s, ASSETS / "public-demo-qr.png", 10.27, 4.70, 1.55, 1.55)
    pic.click_action.hyperlink.address = DEMO

    # Remove only the explicitly excluded instruction slide and content hints.
    for i, slide in enumerate(slides):
        if i == 1:
            continue
        for shape in list(slide.shapes):
            if shape.has_text_frame:
                if shape.text.startswith("Hint (delete):"):
                    remove_shape(shape)
                elif shape.text == "[TEAM NAME]":
                    write(shape, TEAM)
    slide_id = r.slides._sldIdLst[1]
    r.part.drop_rel(slide_id.rId)
    r.slides._sldIdLst.remove(slide_id)
    assert len(r.slides) == 9

    start = 0
    notes_md = [f"# SafeX AI — Code Carnival 3.0 presenter guide\n\n"
                f"Team: **{TEAM}** · ID: **{TEAM_ID}**  \n"
                "Leader: **Divya Gohil** · Member: **Krupa**  \n"
                "Presentation: **10 minutes, nine template slides**  \n"
                "Prepared: **10 October 2026**\n\n"
                "Use Presenter View for the notes embedded in the PowerPoint. "
                "The organizer's instruction slide is excluded from the final deck.\n"]
    titles = []
    for i, s in enumerate(r.slides):
        # Section title is retained byte-for-byte in its visible text.
        title = "CODE CARNIVAL 3.0" if i == 0 else next(
            sh.text for sh in s.shapes if sh.has_text_frame and sh.is_placeholder)
        titles.append(title)
        end = start + TIMES[i]
        timestamp = f"{start//60:02d}:{start%60:02d}–{end//60:02d}:{end%60:02d}"
        note = (f"SLIDE {i+1}: {title}\nPresenter: {SPEAKERS[i]}\n"
                f"Timing: {timestamp} ({TIMES[i]} seconds)\n\n"
                f"SUGGESTED TALK TRACK\n{SCRIPT[i]}\n\n"
                f"SUPPORTING DETAIL / CLAIM BOUNDARIES\n{DETAILS[i]}\n\n"
                "SOURCES / CREDITS\n" + "\n".join(SOURCES[i]))
        s.notes_slide.notes_text_frame.text = note
        notes_md.append(f"## {i+1}. {title}\n\n"
                        f"**{timestamp} · {SPEAKERS[i]} · {TIMES[i]} seconds**\n\n"
                        f"{SCRIPT[i]}\n\n"
                        f"**Supporting details:** {DETAILS[i]}\n\n"
                        "**Sources:** " + "; ".join(SOURCES[i]) + "\n")
        start = end
    assert start == 600

    r.core_properties.title = "SafeX AI — Code Carnival 3.0 Final Submission"
    r.core_properties.subject = "On-device threat, phishing and scam detection"
    r.core_properties.author = "Divya Gohil; Krupa — Team Doclock"
    r.core_properties.keywords = "SafeX AI, Doclock, M87V, offline AI, phishing, Android, Chrome"
    r.save(OUT / f"{STEM}.pptx")
    (OUT / "SafeX-AI-10-Minute-Speaker-Notes.md").write_text("\n\n".join(notes_md), encoding="utf-8")

    # Check final slide sequence, unchanged original images and placeholder removal.
    check = Presentation(OUT / f"{STEM}.pptx")
    original_to_final = {old: new for new, old in enumerate([0, 2, 3, 4, 5, 6, 7, 8, 9])}
    for record in images:
        sh = next(sh for sh in check.slides[original_to_final[record["slide"]]].shapes
                  if sh.name == record["name"] and sh.shape_type == 13)
        assert sha256(sh.image.blob).hexdigest() == record["sha256"]
        assert [sh.left, sh.top, sh.width, sh.height] == record["geometry"]
    for s in check.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                assert not re.search(r"\[[^\]]*\]|Hint \(delete\)|Delete this box", sh.text), sh.text
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        assert "[" not in cell.text
    info = {
        "deck": f"{STEM}.pptx", "pdf": f"{STEM}.pdf", "team": TEAM,
        "teamId": TEAM_ID, "leader": "Divya Gohil", "member": "Krupa",
        "slides": len(check.slides), "durationSeconds": start,
        "slideTitles": titles, "timings": TIMES,
        "templateSha256": sha256(TEMPLATE.read_bytes()).hexdigest(),
        "templateBrandPicturesPreserved": len(images),
        "instructionSlideRemoved": True, "unfilledVisiblePlaceholders": 0,
        "repository": REPO, "repositoryPublicVerified": True,
        "demo": DEMO, "demoHttp200Verified": True,
        "sourceSyncPending": True, "allMemberCommitsVerified": False,
        "themeLabel": "On-device AI Security (descriptive, based on selected problem)",
        "originalSectionTitlesPreserved": titles == original_titles,
        "credits": {"template": "ADSC / Atmiya University, supplied by the user",
                    "logo": "Original SafeX AI SVG from branding/safex-logo.svg",
                    "screenshots": "Existing project screenshots; prior-release captures, cropped",
                    "fonts": "Public Sans and Oswald, SIL OFL; licenses in assets/fonts",
                    "models": "Attribution in models/FRAUD-DATA-ATTRIBUTION.md"},
    }
    assert info["originalSectionTitlesPreserved"], (titles, original_titles)
    (OUT / "submission-validation.json").write_text(json.dumps(info, indent=2)+"\n")
    print(f"Created {OUT / (STEM + '.pptx')}: nine slides, 600 seconds, original branding retained.")


if __name__ == "__main__":
    main()
